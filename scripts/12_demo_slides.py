"""Build the 2-slide backup demo deck (python-pptx, text boxes + images only).

Output: slides/ridescore_crash_risk_demo.pptx
Also writes a layout-check preview (PIL mock render with DejaVu Sans, which is
~10% wider than Arial, so it is a conservative fit test) to
slides/preview/slide{1,2}_mock.png. The mock is NOT a real PowerPoint render.
"""
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "slides" / "ridescore_crash_risk_demo.pptx"
PREVIEW_DIR = ROOT / "slides" / "preview"
BNA_IMG = ROOT / "output" / "bna_compare.png"
MAP_IMG = ROOT / "output" / "lts_disagreement_map.png"

W_IN, H_IN = 13.333, 7.5
INK = "1f2937"
MUTED = "4b5563"
ACCENT = "C2410C"
FONT = "Arial"
M = 0.5  # outer margin (in), >= 0.4
BODY = 17

# ---- paragraph spec: dict(runs=[(text, opts)], size, bullet, before) ---------
# run opts: b (bold), a (accent colour), size (override), c (colour hex)


def P(*runs, size=BODY, bullet=False, before=0, color=INK):
    rr = []
    for r in runs:
        if isinstance(r, str):
            rr.append((r, {}))
        else:
            rr.append(r)
    return dict(runs=rr, size=size, bullet=bullet, before=before, color=color)


def B(t, **kw):
    return (t, dict(b=True, **kw))


def A(t, **kw):  # accent number
    return (t, dict(b=True, a=True, **kw))


SLIDES = []

# ============================ SLIDE 1 =========================================
img_w1 = 6.33
img_h1 = img_w1 * 630 / 1080
s1 = dict(
    boxes=[
        (M, 0.4, W_IN - 2 * M, 0.65, [P(B("Calm on the map, crashes on the street"), size=32)]),
        (M, 1.08, W_IN - 2 * M, 0.7, [P(
            "Checking RideScore DC's LTS & BNA against 5 years of DC bike crashes "
            "· Challenge 1 (ideas 4.3 + 4.2)", size=16)]),
        (M, 1.8, W_IN - 2 * M, 0.35, [P("Yuvraj Gupta · github.com/YuvrajGupta1808/CivicTech",
                                         size=14, color=MUTED)]),
        (M, 2.35, 5.6, 4.7, [
            P(B("Problem: "),
              "RideScore colours DC streets by design (LTS). Nobody had checked those "
              "colours against where cyclists actually crash."),
            P(B("What we did: "),
              "Joined 3,089 bike crashes (Crashes in DC, 2021–2026) to 19,554 DDOT blocks. "
              "Trained a crash model on OSM + DDOT street design (OSM↔DDOT cross-attention "
              "+ boosting + safety-performance formula). Tested only on wards it never saw.",
              before=12),
            P(B("Finding: "), "the 10% of blocks our model ranks riskiest hold ",
              A("44%", size=24), " of crashes. LTS: ", B("13%"), ". BNA: ", B("18%"),
              ". Random: ", B("10%"), ".", before=12),
        ]),
        (6.5, 2.35 + img_h1 + 0.1, img_w1, 0.4, [P(
            "Crashes per km: 'comfortable' vs 'uncomfortable' streets", size=14, color=MUTED)]),
    ],
    images=[(BNA_IMG, 6.5, 2.35, img_w1, img_h1)],
)
SLIDES.append(s1)

# ============================ SLIDE 2 =========================================
img_h2 = 5.35
img_w2 = img_h2 * 823 / 1076
rx = M + img_w2 + 0.45
s2 = dict(
    boxes=[
        (M, 0.4, W_IN - 2 * M, 0.65, [P(B("Where LTS says calm, crashes say otherwise"), size=32)]),
        (rx, 1.2, W_IN - M - rx, 5.4, [
            P(B("Evidence")),
            P("Streets LTS calls calm but our model flags: ", A("4%"), " of the network, ",
              A("15%"), " of all bike crashes (5.3 crashes/km, same as streets both call risky).",
              bullet=True, before=3),
            P("Worst block in DC: 14th St NW at Irving St, LTS 1 with a protected lane, ",
              A("18 crashes"), " in 5 years.", bullet=True, before=3),
            P(B("Honest limits"), before=12),
            P("No ridership data: this shows where crashes happen, not danger per ride "
              "(busy bike routes get more crashes). Wards 7–8 lose ~20% of crashes in the "
              "data match.", before=3),
            P(B("Next steps for RideScore"), before=12),
            P("Add a crash layer next to LTS, not instead of it.", bullet=True, before=3),
            P("Add bike counts to separate busy from dangerous.", bullet=True, before=3),
            P("Fix Ward 7–8 crash matching.", bullet=True, before=3),
            P("Per-segment scores ready: crash_risk_by_segment.parquet.", bullet=True, before=3),
        ]),
        (M, 6.7, W_IN - 2 * M, 0.35, [P(
            "Data: © OpenStreetMap contributors (ODbL); DDOT, MPD Crashes in DC (CC BY 4.0)",
            size=14, color=MUTED)]),
    ],
    images=[(MAP_IMG, M, 1.2, img_w2, img_h2)],
)
SLIDES.append(s2)


# ============================ pptx build ======================================
def rgb(h):
    return RGBColor.from_string(h.upper())


def build():
    prs = Presentation()
    prs.slide_width = Inches(W_IN)
    prs.slide_height = Inches(H_IN)
    blank = prs.slide_layouts[6]
    for spec in SLIDES:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb("FFFFFF")
        for path, x, y, w, h in spec["images"]:
            slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        for x, y, w, h, paras in spec["boxes"]:
            tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.NONE
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            for i, para in enumerate(paras):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                if para["before"]:
                    p.space_before = Pt(para["before"])
                if para["bullet"]:
                    pPr = p._p.get_or_add_pPr()
                    pPr.set("marL", str(Inches(0.25)))
                    pPr.set("indent", str(-Inches(0.25)))
                    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
                    pPr.append(bu)
                for text, o in para["runs"]:
                    r = p.add_run()
                    r.text = text
                    f = r.font
                    f.name = FONT
                    f.size = Pt(o.get("size", para["size"]))
                    f.bold = bool(o.get("b"))
                    f.color.rgb = rgb(ACCENT if o.get("a") else o.get("c", para["color"]))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print("saved", OUT)


# ============================ mock preview / fit check ========================
FONT_DIR = Path(__import__("matplotlib").__file__).parent / "mpl-data" / "fonts" / "ttf"
PPI = 100  # px per inch in the mock


def pil_font(size_pt, bold):
    f = FONT_DIR / ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf")
    return ImageFont.truetype(str(f), max(1, round(size_pt / 72 * PPI)))


def layout_box(draw, x, y, w, h, paras, slide_no):
    """Greedy word wrap with per-run fonts; returns bottom y (in) of text."""
    cur = y * PPI
    for para in paras:
        cur += para["before"] / 72 * PPI
        indent = 0.25 * PPI if para["bullet"] else 0
        maxw = w * PPI - indent
        # tokens: (word_with_trailing_space, font, colour, size)
        toks = []
        for text, o in para["runs"]:
            size = o.get("size", para["size"])
            font = pil_font(size, o.get("b"))
            col = "#" + (ACCENT if o.get("a") else o.get("c", para["color"]))
            parts = text.split(" ")
            for k, wd in enumerate(parts):
                tail = " " if k < len(parts) - 1 else ""
                if wd == "" and tail == "":
                    continue
                toks.append((wd + tail, font, col, size))
        lines, line, lw = [], [], 0
        for t in toks:
            tw = t[1].getlength(t[0].rstrip(" "))
            full = t[1].getlength(t[0])
            if line and lw + tw > maxw:
                lines.append(line)
                line, lw = [], 0
            line.append(t)
            lw += full
        if line:
            lines.append(line)
        for ln in lines:
            sz = max(t[3] for t in ln)
            lh = sz * 1.2 / 72 * PPI
            if para["bullet"] and ln is lines[0]:
                draw.text((x * PPI + 4, cur + (lh - sz / 72 * PPI) * 0.3), "•",
                          font=pil_font(para["size"], False), fill="#" + INK)
            xx = x * PPI + indent
            base = cur + (lh - sz / 72 * PPI) * 0.3
            for wd, font, col, s in ln:
                draw.text((xx, base + (sz - s) / 72 * PPI * 0.8), wd, font=font, fill=col)
                xx += font.getlength(wd)
            cur += lh
    bottom = cur / PPI
    if bottom > y + h + 1e-6:
        print(f"  !! slide {slide_no}: text overflows box at y={y}: bottom {bottom:.2f} > {y + h:.2f}")
    return bottom


def preview():
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    for n, spec in enumerate(SLIDES, 1):
        im = Image.new("RGB", (round(W_IN * PPI), round(H_IN * PPI)), "white")
        d = ImageDraw.Draw(im)
        for path, x, y, w, h in spec["images"]:
            pic = Image.open(path).convert("RGB").resize((round(w * PPI), round(h * PPI)))
            im.paste(pic, (round(x * PPI), round(y * PPI)))
        rects = []
        for x, y, w, h, paras in spec["boxes"]:
            bottom = layout_box(d, x, y, w, h, paras, n)
            print(f"slide {n}: box y={y:.2f}-{y + h:.2f} text bottom {bottom:.2f} "
                  f"(slack {y + h - bottom:+.2f} in)")
            rects.append((x, y, w, max(h, 0) if bottom <= y + h else bottom - y))
        # overlap + margin checks (box rects and images)
        items = [(x, y, w, h) for _, x, y, w, h in spec["images"]] + [
            (x, y, w, h) for x, y, w, h, _ in spec["boxes"]]
        for i, a in enumerate(items):
            if a[0] < 0.4 or a[1] < 0.4 or a[0] + a[2] > W_IN - 0.4 or a[1] + a[3] > H_IN - 0.4:
                print(f"  !! slide {n}: item {a} violates 0.4in margin")
            for b in items[i + 1:]:
                if a[0] < b[0] + b[2] and b[0] < a[0] + a[2] and a[1] < b[1] + b[3] and b[1] < a[1] + a[3]:
                    print(f"  !! slide {n}: overlap {a} vs {b}")
        im.save(PREVIEW_DIR / f"slide{n}_mock.png")


if __name__ == "__main__":
    build()
    preview()
